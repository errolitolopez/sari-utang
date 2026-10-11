"""Reading and writing Excel .xlsx files for the products list (openpyxl)."""

import io

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

from sari_utang.domain.product_import import COLUMNS

EXTRA_DROPDOWN_ROWS = 500  # the unit dropdown also works for rows added later


def products_workbook(products, units):
    """The products as an Excel file in bytes: name, price (in pesos), unit.

    The unit column has a dropdown with the units (listed on a second sheet).
    """
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in products:
        sheet.append([row["name"], row["price"] / 100, row["unit"]])
    last_row = max(sheet.max_row, 2)
    for (price_cell,) in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        price_cell.number_format = "#,##0.00"
    sheet.freeze_panes = "A2"  # header stays visible when scrolling
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12

    unit_sheet = workbook.create_sheet("Units")
    for unit in units:
        unit_sheet.append([unit])
    dropdown = DataValidation(
        type="list", formula1=f"Units!$A$1:$A${len(units)}", allow_blank=False,
        showErrorMessage=True, errorTitle="Unit", error="Choose a unit from the list.",
    )
    sheet.add_data_validation(dropdown)
    dropdown.add(f"C2:C{last_row + EXTRA_DROPDOWN_ROWS}")

    data = io.BytesIO()
    workbook.save(data)
    return data.getvalue()


def read_products_sheet(raw_bytes):
    """The rows of the uploaded file (the "Products" sheet, or the first one), as lists of cells.

    Raises ValueError when the bytes are not an .xlsx file.
    """
    try:
        # read_only: reads row by row; data_only: formulas give their result
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception as error:
        raise ValueError("not an xlsx file") from error
    try:
        sheet = workbook["Products"] if "Products" in workbook.sheetnames else workbook.worksheets[0]
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    finally:
        workbook.close()
