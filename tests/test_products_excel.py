"""Products import and export as Excel .xlsx."""

import io

from openpyxl import Workbook, load_workbook

from conftest import add_product


def excel_file(rows, header=("name", "price", "unit")):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(list(header))
    for row in rows:
        sheet.append(list(row))
    data = io.BytesIO()
    workbook.save(data)
    data.seek(0)
    return data


def upload(client, data, name="products.xlsx"):
    return client.post("/api/products/import", data={"file": (data, name)},
                       content_type="multipart/form-data")


def test_export_has_products_and_unit_dropdown(client):
    add_product(client, "Coke Mismo", 2050, "bottle")
    response = client.get("/products/export.xlsx")
    assert response.status_code == 200
    workbook = load_workbook(io.BytesIO(response.data))
    rows = list(workbook["Products"].iter_rows(values_only=True))
    assert rows == [("name", "price", "unit"), ("Coke Mismo", 20.5, "bottle")]
    assert workbook["Units"]["A1"].value == "pc"
    assert len(workbook["Products"].data_validations.dataValidation) == 1


def test_import_adds_updates_and_leaves_the_rest(client):
    add_product(client, "Milo", 1200, "sachet")
    add_product(client, "Kopiko", 1000, "sachet")
    response = upload(client, excel_file([
        ("Milo", "₱13.50", "SACHET"),   # new price, unit in capitals, peso text
        ("Kopiko", 10, "sachet"),        # same price
        ("Asin", 15.5, "pack"),          # new
    ]))
    assert response.get_json() == {"added": 1, "updated": 1, "unchanged": 1}
    products = {p["name"]: p["price"] for p in client.get("/api/products").get_json()["products"]}
    assert products == {"Milo": 1350, "Kopiko": 1000, "Asin": 1550}


def test_import_is_all_or_nothing(client):
    response = upload(client, excel_file([
        ("Asin", 15, "pack"),
        ("", 10, "pc"),
        ("Itlog", "abc", "pc"),
        ("Mantika", 20, "barrel"),
        ("Asin", 16, "pack"),
    ]))
    data = response.get_json()
    assert response.status_code == 400
    assert data["error_count"] == 4
    assert data["errors"][0].startswith("Row 3: name is required")
    assert "Row 4: price" in data["errors"][1]
    assert "Row 5: unit" in data["errors"][2]
    assert "Row 6: Asin per pack is already in the file above" in data["errors"][3]
    assert client.get("/api/products").get_json()["products"] == []


def test_import_file_problems(client):
    assert upload(client, io.BytesIO(b"not excel")).get_json()["errors"][0].startswith("This is not an Excel")
    wrong_header = upload(client, excel_file([], header=("product", "cost", "unit")))
    assert wrong_header.get_json()["errors"] == ["The first row must be: name, price, unit"]
    assert upload(client, excel_file([])).get_json()["errors"] == ["The file has no products."]
    assert client.post("/api/products/import").get_json()["errors"] == ["Choose an Excel (.xlsx) file."]
