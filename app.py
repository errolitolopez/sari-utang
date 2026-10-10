"""sari-utang: pautang list for a sari-sari store.

Run with: flask --app app run --debug
"""

import io
import json
import os
import re
import secrets
import time
import zipfile
from datetime import date, timedelta
from decimal import ROUND_HALF_UP, Decimal
from urllib.parse import quote
from functools import wraps

from flask import (
    Flask, abort, jsonify, redirect, render_template, request, send_file, session, url_for,
)
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation
from werkzeug.security import check_password_hash

import db
import seed
from invoice_pdf import invoice_pdf_bytes, peso, qty

app = Flask(__name__, instance_relative_config=True)

# The instance/ folder holds private files (database, secret key). It is not committed to git.
os.makedirs(app.instance_path, exist_ok=True)
app.config["DATABASE"] = os.path.join(app.instance_path, "sari_utang.db")

# Stay logged in for 30 days so the PIN isn't asked every time.
app.config["PERMANENT_SESSION_LIFETIME"] = timedelta(days=30)
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"
# Largest upload accepted (e.g. a products Excel file). Bigger requests get "413 Too Large".
app.config["MAX_CONTENT_LENGTH"] = 2 * 1024 * 1024  # 2 MB


def load_secret_key():
    """The secret key signs the login cookie. Create one the first time and reuse it after."""
    path = os.path.join(app.instance_path, "secret_key")
    if not os.path.exists(path):
        with open(path, "w") as f:
            f.write(secrets.token_hex(32))
    with open(path) as f:
        return f.read().strip()


app.secret_key = load_secret_key()
db.init_app(app)
app.cli.add_command(seed.seed_command)


# --- Login protection --------------------------------------------------------

# After too many wrong PINs, block that device for a while so the PIN can't be guessed.
MAX_ATTEMPTS = 5
LOCK_SECONDS = 5 * 60
failed_logins = {}  # ip address -> {"count": wrong tries, "locked_until": timestamp}


def is_logged_in():
    return session.get("logged_in") is True


def login_required_page(view):
    """For HTML pages: send the user to the login page if not logged in."""
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not is_logged_in():
            return redirect(url_for("login_page"))
        return view(*args, **kwargs)
    return wrapper


def login_required_api(view):
    """For API routes: answer 401 if not logged in."""
    @wraps(view)
    def wrapper(*args, **kwargs):
        if not is_logged_in():
            return jsonify(error="Please log in."), 401
        return view(*args, **kwargs)
    return wrapper


# --- Pages -------------------------------------------------------------------

@app.get("/")
def index():
    return redirect(url_for("dashboard_page" if is_logged_in() else "login_page"))


@app.get("/login")
def login_page():
    if is_logged_in():
        return redirect(url_for("dashboard_page"))
    return render_template("login.html")


@app.get("/dashboard")
@login_required_page
def dashboard_page():
    return render_template("dashboard.html")


@app.get("/customers")
@login_required_page
def customers_page():
    return render_template("customers.html")


@app.get("/customers/<int:customer_id>")
@login_required_page
def customer_page(customer_id):
    if find_customer(customer_id) is None:
        abort(404)
    return render_template("customer.html", customer_id=customer_id)


# --- API ---------------------------------------------------------------------

@app.post("/api/login")
def api_login():
    ip = request.remote_addr
    record = failed_logins.get(ip, {"count": 0, "locked_until": 0})

    wait = int(record["locked_until"] - time.time())
    if wait > 0:
        return jsonify(error=f"Too many wrong PINs. Try again in {wait // 60 + 1} min."), 429

    pin_hash = db.get_setting("pin_hash")
    if pin_hash is None:
        return jsonify(error="No PIN set yet. Run: flask --app app set-pin"), 503

    pin = str((request.get_json(silent=True) or {}).get("pin", ""))
    if not check_password_hash(pin_hash, pin):
        record["count"] += 1
        if record["count"] >= MAX_ATTEMPTS:
            record = {"count": 0, "locked_until": time.time() + LOCK_SECONDS}
        failed_logins[ip] = record
        return jsonify(error="Wrong PIN."), 401

    failed_logins.pop(ip, None)
    session.clear()
    session.permanent = True
    session["logged_in"] = True
    return jsonify(ok=True)


@app.post("/api/logout")
def api_logout():
    session.clear()
    return jsonify(ok=True)


# Voided transactions are left out of every balance.
BALANCES_SQL = """
    SELECT c.id, c.name, c.phone,
           COALESCE(SUM(CASE WHEN t.type = 'utang' THEN t.amount ELSE -t.amount END), 0)
               AS balance
    FROM customers c
    LEFT JOIN transactions t ON t.customer_id = c.id AND t.voided_at IS NULL
"""


def customer_balances(order_by):
    """All customers with their balance (total utang - total payments, in centavos)."""
    return db.get_db().execute(f"{BALANCES_SQL} GROUP BY c.id ORDER BY {order_by}").fetchall()


def find_customer(customer_id):
    """One customer with their balance, or None if there is no such customer."""
    return db.get_db().execute(
        f"{BALANCES_SQL} WHERE c.id = ? GROUP BY c.id", (customer_id,)
    ).fetchone()


# --- Singilin (customers to collect from) -------------------------------------

SETTINGS_FILE = os.path.join(app.root_path, "config", "settings.json")


def load_settings():
    """Read config/settings.json. Read every time, so edits show up on the next refresh.

    overdue_days: an utang unpaid for this many days or more counts as "matagal".
    store_name:   shown at the top of invoices (optional).
    """
    with open(SETTINGS_FILE, encoding="utf8") as f:
        settings = json.load(f)
    days = settings.get("overdue_days") if isinstance(settings, dict) else None
    if not isinstance(days, int) or isinstance(days, bool) or not 1 <= days <= 365:
        raise ValueError(f'{SETTINGS_FILE} must look like {{"overdue_days": 30}} (1 to 365)')
    store_name = settings.setdefault("store_name", "Sari-Sari Store")
    if not isinstance(store_name, str) or not 0 < len(store_name.strip()) <= 60:
        raise ValueError(f'{SETTINGS_FILE}: "store_name" must be text, 1 to 60 characters')
    return settings


def collection_list():
    """Everyone who owes, with how long their oldest unpaid utang has been waiting.

    Payments pay off the oldest utang first. So for each customer we add up all
    their payments, then go through their utang from oldest to newest, crossing
    off each one the payments fully cover. The first one left is the oldest unpaid.
    """
    rows = db.get_db().execute(
        """
        SELECT t.customer_id, t.type, t.amount, t.date, c.name, c.phone
        FROM transactions t
        JOIN customers c ON c.id = t.customer_id
        WHERE t.voided_at IS NULL
        ORDER BY t.date, t.id
        """
    ).fetchall()

    customers = {}  # customer id -> what we know about them
    for row in rows:
        c = customers.setdefault(row["customer_id"], {
            "id": row["customer_id"], "name": row["name"], "phone": row["phone"],
            "utang": [], "paid": 0, "last_payment": None,
        })
        if row["type"] == "utang":
            c["utang"].append(row)
        else:
            c["paid"] += row["amount"]
            c["last_payment"] = row["date"]  # rows are oldest first, so the last one wins

    today = date.today()
    result = []
    for c in customers.values():
        balance = sum(u["amount"] for u in c["utang"]) - c["paid"]
        if balance <= 0:
            continue

        left_to_cover = c["paid"]
        unpaid_since = None
        for u in c["utang"]:
            if left_to_cover >= u["amount"]:
                left_to_cover -= u["amount"]  # fully paid
            else:
                unpaid_since = u["date"]  # oldest utang not fully paid yet
                break

        since_day = date.fromisoformat(unpaid_since[:10])  # "2026-10-10 10:15:22" -> 2026-10-10
        result.append({
            "id": c["id"],
            "name": c["name"],
            "phone": c["phone"],
            "balance": balance,
            "unpaid_since": unpaid_since,
            "days": (today - since_day).days,
            "last_payment": c["last_payment"],
        })

    result.sort(key=lambda c: (-c["days"], -c["balance"]))  # longest waiting first
    return result


@app.get("/singilin")
@login_required_page
def singilin_page():
    return render_template("singilin.html")


# Not called /api/collect on purpose: ad/tracker blockers block URLs with "collect"
# (analytics tools use it), which made the page fail with "Cannot reach the server".
@app.get("/api/singilin")
@login_required_api
def api_singilin():
    return jsonify(
        customers=collection_list(),
        overdue_days=load_settings()["overdue_days"],
    )


@app.get("/api/dashboard")
@login_required_api
def api_dashboard():
    conn = db.get_db()

    balances = customer_balances(order_by="balance DESC")
    with_utang = [row for row in balances if row["balance"] > 0]

    recent = conn.execute(
        """
        SELECT t.id, t.type, t.amount, t.note, t.date, t.customer_id, c.name AS customer_name
        FROM transactions t
        JOIN customers c ON c.id = t.customer_id
        WHERE t.voided_at IS NULL
        ORDER BY t.date DESC, t.id DESC
        LIMIT 5
        """
    ).fetchall()

    overdue_days = load_settings()["overdue_days"]
    overdue = [c for c in collection_list() if c["days"] >= overdue_days]

    return jsonify(
        overdue_days=overdue_days,
        overdue_count=len(overdue),
        total_outstanding=sum(row["balance"] for row in with_utang),
        customer_count=len(balances),
        customers_with_utang=len(with_utang),
        top_balances=[dict(row) for row in with_utang[:5]],
        recent_transactions=[dict(row) for row in recent],
    )


@app.get("/api/customers")
@login_required_api
def api_list_customers():
    rows = customer_balances(order_by="c.name COLLATE NOCASE")
    return jsonify(customers=[dict(row) for row in rows])


@app.post("/api/customers")
@login_required_api
def api_add_customer():
    data = request.get_json(silent=True) or {}
    # " ".join(split()) trims the ends and turns double spaces into one.
    name = " ".join(str(data.get("name", "")).split())
    phone = str(data.get("phone") or "").strip()

    if not name:
        return jsonify(error="Name is required."), 400
    if len(name) > 100:
        return jsonify(error="Name is too long (max 100 characters)."), 400
    if phone and not re.fullmatch(r"[0-9+\- ]{7,20}", phone):
        return jsonify(error="Phone can only have numbers, spaces, + and - (7 to 20 characters)."), 400

    conn = db.get_db()
    # Two customers with the same name would be confusing in the list.
    same_name = conn.execute(
        "SELECT 1 FROM customers WHERE name = ? COLLATE NOCASE", (name,)
    ).fetchone()
    if same_name:
        return jsonify(error=f"{name} is already in the list."), 409

    cursor = conn.execute(
        "INSERT INTO customers (name, phone) VALUES (?, ?)", (name, phone or None)
    )
    conn.commit()
    customer = {"id": cursor.lastrowid, "name": name, "phone": phone or None, "balance": 0}
    return jsonify(customer=customer), 201


# --- Products ----------------------------------------------------------------

MAX_PRICE = 100_000_00  # P100,000.00 in centavos, catches typos like an extra zero

UNITS_FILE = os.path.join(app.root_path, "config", "units.json")


def load_units():
    """Read the units (what one "price" buys) from config/units.json.

    The file is read every time, so edits show up on the next page refresh.
    Format: {"kg": "kg (kilo)", ...} -> the key is saved with the product,
    the text is shown in the dropdown. The first unit is the default.
    """
    with open(UNITS_FILE, encoding="utf8") as f:
        units = json.load(f)
    valid = (
        isinstance(units, dict)
        and units
        and all(isinstance(k, str) and k.strip() and isinstance(v, str) for k, v in units.items())
    )
    if not valid:
        raise ValueError(f'{UNITS_FILE} must look like {{"pc": "pc (piraso)", "kg": "kg (kilo)"}}')
    return units


def read_product_form(current_unit=None):
    """Check the name, price and unit sent by the page. Returns (product, error).

    current_unit: when editing, the product's saved unit is still allowed even if
    it was removed from config/units.json, so the product can be saved unchanged.
    """
    data = request.get_json(silent=True) or {}
    name = " ".join(str(data.get("name", "")).split())
    price = data.get("price")
    unit = data.get("unit")

    if not name:
        return None, "Name is required."
    if len(name) > 100:
        return None, "Name is too long (max 100 characters)."
    # bool is a kind of int in Python, so rule it out (True would count as 1).
    if not isinstance(price, int) or isinstance(price, bool) or price <= 0:
        return None, "Price must be more than P0."
    if price > MAX_PRICE:
        return None, "Price is too high (max P100,000)."
    if not isinstance(unit, str) or (unit not in load_units() and unit != current_unit):
        return None, "Choose a unit from the list."
    return {"name": name, "price": price, "unit": unit}, None


def product_taken(name, unit, except_id=None):
    """True if another product already has this name AND unit.

    The same name with a different unit is allowed, e.g. Coke Mismo per pc and per box.
    except_id: when editing, the product being edited doesn't count as "another".
    """
    row = db.get_db().execute(
        "SELECT 1 FROM products WHERE name = ? COLLATE NOCASE AND unit = ? AND id != ?",
        (name, unit, except_id or 0),
    ).fetchone()
    return row is not None


@app.get("/products")
@login_required_page
def products_page():
    return render_template("products.html", units=load_units())


MAX_PER_PAGE = 100


@app.get("/api/products")
@login_required_api
def api_list_products():
    """Products, A-Z.

    Without ?page= : all products (the utang form searches all of them).
    With ?page=2&per_page=20&q=coke : one page of the products whose name has "coke".
    """
    conn = db.get_db()
    page = request.args.get("page", type=int)
    if page is None:
        rows = conn.execute(
            "SELECT id, name, price, unit FROM products ORDER BY name COLLATE NOCASE, price"
        ).fetchall()
        return jsonify(products=[dict(row) for row in rows])

    per_page = min(max(request.args.get("per_page", default=20, type=int), 1), MAX_PER_PAGE)
    query = " ".join(request.args.get("q", "").split())
    # In LIKE, % and _ are wildcards; escape them so searching "50%" finds "50%".
    pattern = "%" + query.replace("\\", "\\\\").replace("%", "\\%").replace("_", "\\_") + "%"
    where = "WHERE name LIKE ? ESCAPE '\\'"

    total = conn.execute(f"SELECT COUNT(*) FROM products {where}", (pattern,)).fetchone()[0]
    pages = max((total + per_page - 1) // per_page, 1)
    page = min(max(page, 1), pages)  # e.g. page 9 of 3 -> page 3
    rows = conn.execute(
        f"SELECT id, name, price, unit FROM products {where} "
        "ORDER BY name COLLATE NOCASE, price LIMIT ? OFFSET ?",
        (pattern, per_page, (page - 1) * per_page),
    ).fetchall()
    all_count = conn.execute("SELECT COUNT(*) FROM products").fetchone()[0]
    return jsonify(
        products=[dict(row) for row in rows],
        page=page, per_page=per_page, pages=pages, total=total, all_count=all_count,
    )


@app.post("/api/products")
@login_required_api
def api_add_product():
    product, error = read_product_form()
    if error:
        return jsonify(error=error), 400
    if product_taken(product["name"], product["unit"]):
        return jsonify(error=f"{product['name']} per {product['unit']} is already in the list."), 409

    conn = db.get_db()
    cursor = conn.execute(
        "INSERT INTO products (name, price, unit) VALUES (:name, :price, :unit)", product
    )
    conn.commit()
    return jsonify(product={"id": cursor.lastrowid, **product}), 201


@app.put("/api/products/<int:product_id>")
@login_required_api
def api_update_product(product_id):
    conn = db.get_db()
    saved = conn.execute("SELECT unit FROM products WHERE id = ?", (product_id,)).fetchone()
    if saved is None:
        return jsonify(error="Product not found."), 404

    product, error = read_product_form(current_unit=saved["unit"])
    if error:
        return jsonify(error=error), 400
    if product_taken(product["name"], product["unit"], except_id=product_id):
        return jsonify(error=f"{product['name']} per {product['unit']} is already in the list."), 409

    conn.execute(
        "UPDATE products SET name = :name, price = :price, unit = :unit WHERE id = :id",
        {**product, "id": product_id},
    )
    conn.commit()
    return jsonify(product={"id": product_id, **product})


# --- Customer history and utang ----------------------------------------------

MAX_ITEMS = 50
MAX_QUANTITY = Decimal("10000")
MAX_UTANG_TOTAL = 1_000_000_00  # P1,000,000.00 in centavos


@app.get("/api/customers/<int:customer_id>/transactions")
@login_required_api
def api_customer_transactions(customer_id):
    customer = find_customer(customer_id)
    if customer is None:
        return jsonify(error="Customer not found."), 404

    conn = db.get_db()
    transactions = [
        dict(row, items=[])
        for row in conn.execute(
            """
            SELECT id, type, amount, note, date, voided_at, void_reason, tendered FROM transactions
            WHERE customer_id = ?
            ORDER BY date DESC, id DESC
            """,
            (customer_id,),
        )
    ]

    # Put each item under its transaction.
    by_id = {t["id"]: t for t in transactions}
    items = conn.execute(
        """
        SELECT i.transaction_id, i.product_id, i.name, i.unit, i.price, i.quantity, i.amount
        FROM transaction_items i
        JOIN transactions t ON t.id = i.transaction_id
        WHERE t.customer_id = ?
        ORDER BY i.id
        """,
        (customer_id,),
    )
    for item in items:
        by_id[item["transaction_id"]]["items"].append(dict(item))

    return jsonify(customer=dict(customer), transactions=transactions)


def parse_quantity(value):
    """A quantity like 2 or 0.5 as a Decimal, or None if it isn't valid.

    Up to 3 decimals are allowed (0.125 kg). Decimal is used instead of float so
    0.1 + 0.2 is exactly 0.3 when computing amounts.
    """
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    quantity = Decimal(str(value))
    if not quantity.is_finite() or quantity <= 0 or quantity > MAX_QUANTITY:
        return None
    if quantity.as_tuple().exponent < -3:
        return None
    return quantity


def line_amount(price, quantity):
    """price (centavos) x quantity, rounded to the nearest centavo (.5 rounds up)."""
    return int((price * quantity).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def read_utang_item(raw, number):
    """Check one item of an utang. Returns (item, error).

    An item is either a product from the list: {"product_id": 3, "quantity": 2}
    or typed in by hand:                        {"name": "Yelo", "price": 500, "quantity": 1}
    """
    if not isinstance(raw, dict):
        return None, f"Item {number} is not valid."

    quantity = parse_quantity(raw.get("quantity"))
    if quantity is None:
        return None, f"Item {number}: quantity must be more than 0 (up to 3 decimals)."

    if raw.get("product_id") is not None:
        product_id = raw["product_id"]
        product = None
        if isinstance(product_id, int) and not isinstance(product_id, bool):
            product = db.get_db().execute(
                "SELECT id, name, unit, price FROM products WHERE id = ?", (product_id,)
            ).fetchone()
        if product is None:
            return None, f"Item {number}: product not found. Refresh the page and try again."
        name, unit, price = product["name"], product["unit"], product["price"]
    else:
        product_id, unit = None, None
        name = " ".join(str(raw.get("name", "")).split())
        price = raw.get("price")
        if not name or len(name) > 100:
            return None, f"Item {number}: name is required (max 100 characters)."
        if not isinstance(price, int) or isinstance(price, bool) or not 0 < price <= MAX_PRICE:
            return None, f"Item {number}: price must be more than P0 (max P100,000)."

    amount = line_amount(price, quantity)
    if amount <= 0:
        return None, f"Item {number}: quantity is too small, the amount would be P0."

    item = {
        "product_id": product_id,
        "name": name,
        "unit": unit,
        "price": price,
        "quantity": float(quantity),
        "amount": amount,
    }
    return item, None


@app.post("/api/transactions")
@login_required_api
def api_add_transaction():
    """Add an utang or a payment for a customer."""
    data = request.get_json(silent=True) or {}

    kind = data.get("type")
    if kind not in ("utang", "payment"):
        return jsonify(error='type must be "utang" or "payment".'), 400

    customer_id = data.get("customer_id")
    customer = None
    if isinstance(customer_id, int) and not isinstance(customer_id, bool):
        customer = find_customer(customer_id)
    if customer is None:
        return jsonify(error="Customer not found."), 404

    note = " ".join(str(data.get("note") or "").split())
    if len(note) > 200:
        return jsonify(error="Note is too long (max 200 characters)."), 400

    if kind == "utang":
        return add_utang(customer, data.get("items"), note)
    return add_payment(customer, data.get("amount"), note)


def add_utang(customer, raw_items, note):
    if not isinstance(raw_items, list) or not raw_items:
        return jsonify(error="Add at least one item."), 400
    if len(raw_items) > MAX_ITEMS:
        return jsonify(error=f"Too many items (max {MAX_ITEMS})."), 400

    # Check every item first, so nothing is saved if one of them is wrong.
    items = []
    for number, raw in enumerate(raw_items, start=1):
        item, error = read_utang_item(raw, number)
        if error:
            return jsonify(error=error), 400
        items.append(item)

    # The total is always computed here, never taken from the page.
    total = sum(item["amount"] for item in items)
    if total > MAX_UTANG_TOTAL:
        return jsonify(error="Total is too high (max P1,000,000)."), 400

    conn = db.get_db()
    cursor = conn.execute(
        "INSERT INTO transactions (customer_id, type, amount, note) VALUES (?, 'utang', ?, ?)",
        (customer["id"], total, note or None),
    )
    transaction_id = cursor.lastrowid
    conn.executemany(
        """
        INSERT INTO transaction_items
            (transaction_id, product_id, name, unit, price, quantity, amount)
        VALUES (:transaction_id, :product_id, :name, :unit, :price, :quantity, :amount)
        """,
        [{**item, "transaction_id": transaction_id} for item in items],
    )
    conn.commit()  # the transaction and its items are saved together

    return jsonify(transaction={"id": transaction_id, "amount": total, "items": items}), 201


MAX_TENDERED = 1_000_000_00  # P1,000,000.00 in centavos, catches typos


def add_payment(customer, amount, note):
    """A payment lowers the balance. It can be partial.

    If the customer hands over more than the utang, it counts as paid in full:
    only the utang is recorded as paid, and the extra is the sukli (change).
    """
    if not isinstance(amount, int) or isinstance(amount, bool) or amount <= 0:
        return jsonify(error="Amount must be more than P0."), 400
    if amount > MAX_TENDERED:
        return jsonify(error="Amount is too high (max P1,000,000)."), 400
    if customer["balance"] <= 0:
        return jsonify(error=f"{customer['name']} has no utang to pay."), 400

    paid = min(amount, customer["balance"])
    change = amount - paid  # sukli
    tendered = amount if change > 0 else None  # cash received, kept only when there was sukli

    conn = db.get_db()
    cursor = conn.execute(
        "INSERT INTO transactions (customer_id, type, amount, note, tendered) "
        "VALUES (?, 'payment', ?, ?, ?)",
        (customer["id"], paid, note or None, tendered),
    )
    conn.commit()

    return jsonify(
        transaction={"id": cursor.lastrowid, "amount": paid, "tendered": tendered},
        change=change,
        balance=customer["balance"] - paid,
    ), 201


@app.post("/api/transactions/<int:transaction_id>/void")
@login_required_api
def api_void_transaction(transaction_id):
    """Cancel an utang or payment entered by mistake. It stays in the history, crossed out."""
    conn = db.get_db()
    transaction = conn.execute(
        "SELECT id, customer_id, type, amount, voided_at FROM transactions WHERE id = ?",
        (transaction_id,),
    ).fetchone()
    if transaction is None:
        return jsonify(error="Transaction not found."), 404
    if transaction["voided_at"] is not None:
        return jsonify(error="This is already voided."), 400

    data = request.get_json(silent=True) or {}
    reason = " ".join(str(data.get("reason") or "").split())
    if len(reason) > 200:
        return jsonify(error="Reason is too long (max 200 characters)."), 400

    # Voiding an utang lowers the balance. If part of it was already paid, the
    # balance would go below zero, so the payment has to be voided first.
    customer = find_customer(transaction["customer_id"])
    if transaction["type"] == "utang" and customer["balance"] - transaction["amount"] < 0:
        return jsonify(
            error="Part of this utang was already paid. Void the payment first, then this utang."
        ), 400

    conn.execute(
        "UPDATE transactions SET voided_at = datetime('now', 'localtime'), void_reason = ? "
        "WHERE id = ?",
        (reason or None, transaction_id),
    )
    conn.commit()

    change = -transaction["amount"] if transaction["type"] == "utang" else transaction["amount"]
    return jsonify(ok=True, balance=customer["balance"] + change)


# --- Invoice -----------------------------------------------------------------

# {{ amount|peso }} -> "₱12.50" and {{ item.quantity|qty }} -> "0.5" in templates.
# The functions live in invoice_pdf.py so the page and the PDF format numbers the same way.
peso_filter = app.template_filter("peso")(peso)
quantity_filter = app.template_filter("qty")(qty)


def describe_item(item):
    """"2 bottle Coke Mismo", or "3 × Yelo" for typed-in items (no unit)."""
    quantity = quantity_filter(item["quantity"])
    return f"{quantity} {item['unit']} {item['name']}" if item["unit"] else f"{quantity} × {item['name']}"


def invoice_data(customer_id):
    """Only the utang that is still unpaid.

    Payments pay off the oldest utang first (same rule as the Singilin list):
    - an utang fully covered by payments is left out,
    - an utang partly covered is shown with how much was already paid on it,
    - newer utang are shown in full.
    Voided entries are left out.
    """
    conn = db.get_db()
    rows = conn.execute(
        """
        SELECT id, type, amount, note, date FROM transactions
        WHERE customer_id = ? AND voided_at IS NULL
        ORDER BY date, id
        """,
        (customer_id,),
    ).fetchall()

    payments = [r for r in rows if r["type"] == "payment"]
    paid_left = sum(p["amount"] for p in payments)  # payments not yet applied to an utang

    entries = []
    for row in rows:
        if row["type"] != "utang":
            continue
        if paid_left >= row["amount"]:
            paid_left -= row["amount"]  # fully paid: not on the invoice
            continue
        paid = paid_left  # 0, or part of this utang that was already paid
        paid_left = 0
        entries.append(dict(row, items=[], paid=paid, remaining=row["amount"] - paid))

    by_id = {e["id"]: e for e in entries}
    if by_id:
        placeholders = ",".join("?" * len(by_id))
        for item in conn.execute(
            f"SELECT transaction_id, name, unit, price, quantity, amount FROM transaction_items "
            f"WHERE transaction_id IN ({placeholders}) ORDER BY id",
            list(by_id),
        ):
            by_id[item["transaction_id"]]["items"].append(dict(item))

    # Total quantity and amount per product (same name + unit), e.g. 6 bottle Coke Mismo.
    products = {}
    for entry in entries:
        for item in entry["items"]:
            key = (item["name"], item["unit"])
            total = products.setdefault(
                key, {"name": item["name"], "unit": item["unit"], "quantity": Decimal("0"), "amount": 0}
            )
            total["quantity"] += Decimal(str(item["quantity"]))
            total["amount"] += item["amount"]
    product_totals = sorted(products.values(), key=lambda p: -p["amount"])

    total_utang = sum(e["amount"] for e in entries)
    partly_paid = sum(e["paid"] for e in entries)
    return {
        "entries": entries,
        "products": product_totals,
        "total_utang": total_utang,
        "partly_paid": partly_paid,
        "balance": total_utang - partly_paid,
        "last_payment": payments[-1]["date"] if payments else None,
    }


def invoice_text(store_name, customer, invoice, today):
    """Plain-text version for copying into Messenger/Viber or sending by SMS."""
    lines = [store_name, f"Utang ni {customer['name']} (as of {today})", ""]
    if not invoice["entries"]:
        return "\n".join(lines + ["Walang utang. Bayad na lahat."])
    for entry in invoice["entries"]:
        items = ", ".join(describe_item(i) for i in entry["items"]) or "Utang"
        lines.append(f"{entry['date'][:10]}  {items}: {peso_filter(entry['amount'])}")
        if entry["paid"]:
            lines.append(
                f"   bayad na -{peso_filter(entry['paid'])}, natitira {peso_filter(entry['remaining'])}"
            )
    lines += ["", f"Total utang: {peso_filter(invoice['total_utang'])}"]
    if invoice["partly_paid"]:
        lines.append(f"Bayad na (partial): -{peso_filter(invoice['partly_paid'])}")
    lines.append(f"Balance: {peso_filter(invoice['balance'])}")
    return "\n".join(lines)


@app.get("/customers/<int:customer_id>/invoice")
@login_required_page
def invoice_page(customer_id):
    customer = find_customer(customer_id)
    if customer is None:
        abort(404)
    store_name = load_settings()["store_name"]
    invoice = invoice_data(customer_id)
    today = date.today().isoformat()
    text = invoice_text(store_name, customer, invoice, today)

    sms_link = None
    if customer["phone"]:
        phone = re.sub(r"[^0-9+]", "", customer["phone"])
        sms_link = f"sms:{phone}?body={quote(text)}"

    return render_template(
        "invoice.html",
        store_name=store_name,
        customer=customer,
        invoice=invoice,
        today=today,
        text=text,
        sms_link=sms_link,
    )


# --- Invoice PDFs and batch export (ZIP) -----------------------------------------

def safe_filename(name, fallback):
    """A customer name that is safe as a file name on phones and Windows.

    Removes characters like / \\ : * ? " < > | that file names can't have.
    """
    cleaned = re.sub(r'[\\/:*?"<>|\x00-\x1f]', "", name).strip(" .")
    return cleaned[:80] or fallback


@app.get("/customers/<int:customer_id>/invoice.pdf")
@login_required_page
def invoice_pdf_download(customer_id):
    customer = find_customer(customer_id)
    if customer is None:
        abort(404)
    today = date.today().isoformat()
    pdf = invoice_pdf_bytes(load_settings()["store_name"], customer, invoice_data(customer_id), today)
    filename = safe_filename(customer["name"], f"customer-{customer_id}") + ".pdf"
    return send_file(io.BytesIO(pdf), mimetype="application/pdf", as_attachment=True, download_name=filename)


@app.get("/singilin/invoices.zip")
@login_required_page
def invoices_zip_download():
    """One PDF invoice per customer with utang, named after the customer, in one ZIP.

    ?min_days=30 only includes customers whose oldest unpaid utang is 30+ days old
    (the same filter as the Singilin page). Without it, everyone with utang.
    """
    min_days = request.args.get("min_days", default=0, type=int)
    customers = [c for c in collection_list() if c["days"] >= min_days]
    if not customers:
        abort(404, "No customers with utang for this filter.")

    store_name = load_settings()["store_name"]
    today = date.today().isoformat()
    used_names = set()
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for c in customers:
            customer = find_customer(c["id"])
            pdf = invoice_pdf_bytes(store_name, customer, invoice_data(c["id"]), today)

            # Two names can become the same after cleaning (e.g. "Jun?" and "Jun"): add (2), (3)...
            base = safe_filename(customer["name"], f"customer-{customer['id']}")
            filename, number = f"{base}.pdf", 2
            while filename.lower() in used_names:
                filename, number = f"{base} ({number}).pdf", number + 1
            used_names.add(filename.lower())

            archive.writestr(filename, pdf)

    buffer.seek(0)
    return send_file(
        buffer, mimetype="application/zip", as_attachment=True,
        download_name=f"invoices-{today}.zip",
    )


# --- Products import / export (Excel .xlsx) ------------------------------------

COLUMNS = ["name", "price", "unit"]
MAX_IMPORT_ROWS = 5000


@app.get("/products/export.xlsx")
@login_required_page
def products_export():
    """All products as an Excel file: name, price (in pesos), unit.

    The unit column has a dropdown with the units from config/units.json.
    """
    rows = db.get_db().execute(
        "SELECT name, price, unit FROM products ORDER BY name COLLATE NOCASE, price"
    ).fetchall()
    units = list(load_units())

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Products"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row["name"], row["price"] / 100, row["unit"]])
    last_row = max(sheet.max_row, 2)
    for (price_cell,) in sheet.iter_rows(min_row=2, min_col=2, max_col=2):
        price_cell.number_format = "#,##0.00"
    sheet.freeze_panes = "A2"  # header stays visible when scrolling
    sheet.column_dimensions["A"].width = 32
    sheet.column_dimensions["B"].width = 12
    sheet.column_dimensions["C"].width = 12

    # The units go on a second sheet; the unit cells get a dropdown that lists them.
    unit_sheet = workbook.create_sheet("Units")
    for unit in units:
        unit_sheet.append([unit])
    dropdown = DataValidation(
        type="list", formula1=f"Units!$A$1:$A${len(units)}", allow_blank=False,
        showErrorMessage=True, errorTitle="Unit", error="Choose a unit from the list.",
    )
    sheet.add_data_validation(dropdown)
    dropdown.add(f"C2:C{last_row + 500}")  # also for rows added later

    data = io.BytesIO()
    workbook.save(data)
    data.seek(0)
    return send_file(
        data, as_attachment=True, download_name=f"products-{date.today().isoformat()}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


def parse_import_price(value):
    """A price cell -> centavos, or None if it isn't valid.

    Excel usually gives a number (20.5). A cell typed as text can look like
    "20", "20.50" or "₱1,234.50".
    """
    if isinstance(value, bool) or value is None:
        return None
    if isinstance(value, (int, float)):
        amount = Decimal(str(value))
    else:
        cleaned = str(value).replace("₱", "").replace("P", "").replace(",", "").strip()
        if not re.fullmatch(r"\d+(\.\d+)?", cleaned):
            return None
        amount = Decimal(cleaned)
    if not amount.is_finite() or amount.as_tuple().exponent < -2:
        return None  # more than 2 decimals, e.g. 20.555
    centavos = int(amount * 100)
    return centavos if 0 < centavos <= MAX_PRICE else None


def cell_text(value):
    """A cell as text: None -> "", 555 (a number) -> "555", 12.0 -> "12"."""
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value)


def read_products_excel(raw_bytes):
    """Check every row of an uploaded Excel file. Returns (products, errors).

    Nothing is saved here; if errors is not empty, the import is cancelled.
    """
    try:
        # read_only: reads row by row; data_only: formulas give their result
        workbook = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
    except Exception:
        return [], ["This is not an Excel .xlsx file. In Excel, use Save As → Excel Workbook (.xlsx)."]

    try:
        sheet = workbook["Products"] if "Products" in workbook.sheetnames else workbook.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        header = [cell_text(h).strip().lower() for h in next(rows, ())][:3]
        if header != COLUMNS:
            return [], [f"The first row must be: {', '.join(COLUMNS)}"]

        units = load_units()
        units_lower = {u.lower(): u for u in units}  # "KG" in the file -> "kg"
        products, errors, seen = [], [], set()
        for line_number, row in enumerate(rows, start=2):
            row = (tuple(row) + (None, None, None))[:3]
            if all(cell_text(cell).strip() == "" for cell in row):
                continue  # skip empty rows
            if len(products) + len(errors) >= MAX_IMPORT_ROWS:
                errors.append(f"Too many rows (max {MAX_IMPORT_ROWS}).")
                break
            name = " ".join(cell_text(row[0]).split())
            price = parse_import_price(row[1])
            unit = units_lower.get(cell_text(row[2]).strip().lower())

            problems = []
            if not name or len(name) > 100:
                problems.append("name is required (max 100 characters)")
            if cell_text(row[1]).strip() == "":
                problems.append("price is empty")
            elif price is None:
                problems.append(f'price "{cell_text(row[1])}" is not valid (e.g. 20 or 20.50, max 100,000)')
            if cell_text(row[2]).strip() == "":
                problems.append("unit is empty")
            elif unit is None:
                problems.append(f'unit "{cell_text(row[2])}" is not in config/units.json')
            key = (name.lower(), unit)
            if not problems and key in seen:
                problems.append(f"{name} per {unit} is already in the file above")
            seen.add(key)

            if problems:
                errors.append(f"Row {line_number}: " + "; ".join(problems))
            else:
                products.append({"name": name, "price": price, "unit": unit})
    finally:
        workbook.close()

    if not products and not errors:
        errors.append("The file has no products.")
    return products, errors


@app.post("/api/products/import")
@login_required_api
def api_import_products():
    """Add new products and update prices of existing ones (same name + unit).

    All or nothing: if any row has a problem, nothing is saved.
    Products that are not in the file are left as they are.
    """
    upload = request.files.get("file")
    if upload is None or not upload.filename:
        return jsonify(errors=["Choose an Excel (.xlsx) file."]), 400

    products, errors = read_products_excel(upload.read())
    if errors:
        return jsonify(errors=errors[:20], error_count=len(errors)), 400

    conn = db.get_db()
    added = updated = unchanged = 0
    for p in products:
        existing = conn.execute(
            "SELECT id, price FROM products WHERE name = ? COLLATE NOCASE AND unit = ?",
            (p["name"], p["unit"]),
        ).fetchone()
        if existing is None:
            conn.execute("INSERT INTO products (name, price, unit) VALUES (:name, :price, :unit)", p)
            added += 1
        elif existing["price"] != p["price"]:
            conn.execute("UPDATE products SET price = ? WHERE id = ?", (p["price"], existing["id"]))
            updated += 1
        else:
            unchanged += 1
    conn.commit()  # everything is saved together

    return jsonify(added=added, updated=updated, unchanged=unchanged)
